"""Read a document (PDF, Word, PowerPoint, text) into plain structured text —
the brief for Slides · Presentation (POST /ai/motion/read-doc → the Studio's
prompt box), used like a text the person pasted: headings stay headings
("# …"), bullets stay bullets ("- …"), tables become "a | b | c" rows, a
PowerPoint keeps its slides ("## Slide 3: …").

A long document is condensed to fit the brief (MAX_CHARS): every heading is
kept and the room is shared across the sections, so the deck still covers the
whole outline. A scanned PDF (pictures of pages, no real text) can't be read —
the person is asked to paste the text instead.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path

log = logging.getLogger(__name__)

MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 18000  # the brief takes 20,000 — room for the person's own line on top
KINDS = {".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".txt": "text", ".md": "text"}


class DocError(ValueError):
    """Shown to the person as it is (plain words)."""


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r", "")
    text = re.sub(r"(?<=\S)[ \t ]+", " ", text)  # between words — a line's indent (a sub-bullet) stays
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                raise DocError("This PDF is locked with a password — save it without one and attach it again.") from None
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
    except DocError:
        raise
    except Exception:  # noqa: BLE001 — a damaged file
        log.warning("docread: unreadable PDF", exc_info=True)
        raise DocError("This PDF couldn't be read — try saving it again, or paste its text.") from None
    text = "\n\n".join(p for p in pages if p)
    if len(text) < 40 * max(1, min(len(pages), 5)):
        raise DocError("This PDF is pictures of pages (scanned), so its words can't be read — paste the text instead.")
    return text, len(pages)


def _docx(data: bytes) -> tuple[str, int]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001
        log.warning("docread: unreadable Word file", exc_info=True)
        raise DocError("This Word file couldn't be read — save it as .docx and attach it again.") from None
    out: list[str] = []
    body = d.element.body
    paragraphs = {p._p: p for p in d.paragraphs}
    tables = {t._tbl: t for t in d.tables}
    for child in body.iterchildren():  # in document order: paragraphs and tables
        if child in paragraphs:
            p = paragraphs[child]
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            level = re.search(r"heading (\d)", style)
            if style == "title":
                out.append(f"# {text}")
            elif level:
                out.append(f"{'#' * min(3, int(level[1]) + 1)} {text}" if int(level[1]) > 1 else f"# {text}")
            elif "list" in style:
                out.append(f"- {text}")
            else:
                out.append(text)
        elif child in tables:
            for row in tables[child].rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    out.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat — once each
    return "\n".join(out), 0


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception:  # noqa: BLE001
        log.warning("docread: unreadable PowerPoint", exc_info=True)
        raise DocError("This PowerPoint couldn't be read — save it as .pptx and attach it again.") from None
    out: list[str] = []
    for n, slide in enumerate(deck.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## Slide {n}: {title}" if title else f"## Slide {n}")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        out.append(f"{'  ' * para.level}- {text}")
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(cells):
                        out.append(" | ".join(cells))
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
        if notes:
            out.append(f"(notes: {notes})")
    return "\n".join(out), len(deck.slides)


def condense(text: str, limit: int = MAX_CHARS) -> tuple[str, bool]:
    """``text`` cut to ``limit`` characters without losing its outline: every
    heading stays, and each section keeps its share of the room (its opening lines)."""
    if len(text) <= limit:
        return text, False
    lines = text.split("\n")
    sections: list[list[str]] = [[]]
    for line in lines:
        if re.match(r"#{1,3} ", line) and sections[-1]:
            sections.append([])
        sections[-1].append(line)
    heads = sum(len(s[0]) + 1 for s in sections if s and re.match(r"#{1,3} ", s[0]))
    room = max(limit - heads - 200, limit // 3)
    body = sum(len("\n".join(s)) for s in sections) or 1
    out: list[str] = []
    for s in sections:
        share = max(160, int(room * len("\n".join(s)) / body))
        kept, used = [], 0
        for k, line in enumerate(s):
            if k == 0 and re.match(r"#{1,3} ", line):
                kept.append(line)
                continue
            if used + len(line) > share:
                if used < share * 0.6:  # room for the start of a long paragraph
                    kept.append(line[: max(0, share - used)].rsplit(" ", 1)[0] + " …")
                break
            kept.append(line)
            used += len(line) + 1
        out.append("\n".join(kept))
    joined = "\n".join(out)
    return joined[:limit], True


def read(name: str, data: bytes) -> dict:
    """{"text", "kind", "pages", "condensed", "chars"} for a document — raises DocError."""
    kind = KINDS.get(Path(name or "").suffix.lower())
    if kind is None:
        raise DocError("Attach a PDF, a Word (.docx) or PowerPoint (.pptx) file, or a text file.")
    if len(data) > MAX_BYTES:
        raise DocError("This file is over 20 MB — attach a smaller one.")
    if kind == "pdf":
        text, pages = _pdf(data)
    elif kind == "docx":
        text, pages = _docx(data)
    elif kind == "pptx":
        text, pages = _pptx(data)
    else:
        for enc in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise DocError("This text file couldn't be read — save it as UTF-8.")
        pages = 0
    text = _clean(text)
    if len(text) < 20:
        raise DocError("There's no text in this file to make slides from.")
    full = len(text)
    text, condensed = condense(text)
    return {"text": text, "kind": kind, "pages": pages, "condensed": condensed, "chars": full}
